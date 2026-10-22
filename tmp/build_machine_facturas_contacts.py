from pathlib import Path
from PIL import Image, ImageChops
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, PageBreak, NextPageTemplate, Table, TableStyle, KeepTogether)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

PROJECT = Path(r'C:\Users\Administrador\Desktop\facturacao-angola')
OUT = PROJECT / 'docs'
TMP = PROJECT / 'tmp'
DOCX_PATH = OUT / 'Machine-Facturas-Apresentacao-Institucional.docx'
PDF_PATH = TMP / 'Machine-Facturas-Apresentacao-Institucional-contactos.pdf'
MD_PATH = OUT / 'Machine-Facturas-Post-LinkedIn.md'
LOGO_SOURCE = PROJECT / 'public' / 'logo.png'
LOGO = TMP / 'machine-facturas-logo-trim.png'

OUT.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)

# Crop only the whitespace around the supplied project logo; do not alter the source asset.
logo_image = Image.open(LOGO_SOURCE).convert('RGB')
white = Image.new('RGB', logo_image.size, 'white')
bounds = ImageChops.difference(logo_image, white).getbbox()
if bounds:
    logo_image.crop(bounds).save(LOGO, quality=95)
else:
    logo_image.save(LOGO, quality=95)

sections = [
    ('Apresentação da Machine Facturas', [
        ('p', 'A Machine Facturas é uma plataforma digital de facturação e gestão empresarial. Reúne num espaço de trabalho os registos de clientes e produtos, a preparação e emissão de documentos comerciais, o acompanhamento de pagamentos e a consulta de informação financeira.'),
        ('p', 'O projecto procura apoiar empresas que precisam de manter os seus processos administrativos mais claros e consistentes. Ao relacionar dados comerciais com documentos e movimentos financeiros, a aplicação ajuda a reduzir a dispersão da informação e facilita a consulta do histórico de cada empresa.'),
        ('p', 'A plataforma foi concebida e desenvolvida por Joaquim Zumbuca, numa iniciativa tecnológica própria que alia uma visão prática do software ao compromisso com soluções digitais úteis para as empresas, especialmente no contexto angolano.'),
    ]),
    ('O desafio', [
        ('p', 'Quando os dados de clientes, artigos, facturas e pagamentos ficam repartidos por ficheiros ou ferramentas diferentes, tarefas simples podem exigir pesquisa repetida e reconciliação manual. A falta de uma referência comum também pode dificultar a leitura do estado dos documentos e dos valores por receber.'),
        ('p', 'A necessidade de organização não é igual em todas as empresas. Ainda assim, centralizar registos e manter uma sequência de trabalho coerente pode apoiar a rotina de quem emite documentos, acompanha cobranças ou prepara informação de gestão.'),
    ]),
    ('A solução', [
        ('p', 'A plataforma combina um perfil empresarial com módulos de clientes, produtos, facturação, pagamentos e relatórios. O utilizador pode preparar um rascunho, emitir o documento, consultar o respectivo estado e acompanhar os movimentos associados. A informação é apresentada por empresa, para que cada espaço de trabalho mantenha os seus próprios registos.'),
        ('p', 'A proposta é pragmática: organizar tarefas frequentes num fluxo digital acessível, com valores em kwanzas (Kz) e opções adequadas à actividade comercial local. O produto é desenvolvido de forma incremental e esta apresentação distingue as capacidades presentes no código das áreas que ainda carecem de validação.'),
    ]),
]
features = [
    ('Facturação e notas de crédito', 'A aplicação permite preparar rascunhos e emitir facturas, gerar o documento em PDF, consultar o histórico e emitir notas de crédito associadas a documentos anteriores. Os dados do perfil empresarial e o logótipo podem ser guardados como referência da factura emitida.'),
    ('Clientes e produtos', 'Os registos de clientes incluem identificação e dados de contacto. O catálogo de produtos guarda descrição, preço, taxa de imposto, código, categoria e stock, apoiando a reutilização de informação na preparação dos documentos.'),
    ('Pagamentos e saldos', 'É possível registar pagamentos ligados a facturas, acompanhar estados como pendente, confirmado ou estornado e consultar os montantes pagos e em dívida. O fluxo é de registo e acompanhamento; não foi encontrada integração com gateways de pagamento.'),
    ('Relatórios e exportação', 'A área de relatórios permite consultar dados por período e outros filtros e exportar ficheiros em PDF, Excel (XLSX) ou CSV. O painel apresenta indicadores agregados para acompanhar actividade documental e financeira.'),
    ('Perfil da empresa', 'Cada empresa pode manter dados como denominação, NIF, endereço, contactos, logótipo e contas bancárias. A factura pode guardar uma fotografia dos dados utilizados na emissão, preservando a representação histórica do documento.'),
]
angola_security = [
    ('Pensada para o contexto angolano', [
        ('p', 'O modelo do produto contempla valores em AOA, formatação em português de Angola e datas no fuso horário Africa/Luanda. O perfil empresarial prevê campos como NIF, município e província, e a factura pode apresentar dados bancários e identificação da empresa.'),
        ('p', 'O código inclui uma rota de exportação XML identificada como SAFT-AO, mas a própria implementação descreve a estrutura como simplificada e contém valores de exemplo predefinidos. Esta área deve ser considerada experimental; não é apresentada como exportação fiscal validada, certificação ou integração oficial com a AGT.'),
    ]),
    ('Segurança e organização da informação', [
        ('p', 'O código implementa autenticação por email e palavra-passe, verificação de email e recuperação de palavra-passe. O acesso através do Google OAuth está disponível quando as respectivas credenciais estão configuradas. O acesso às rotas de negócio é validado no servidor, com verificações de sessão, estado da conta, permissões e associação à empresa. Há também limitação de pedidos em rotas de autenticação e registo de acções administrativas.'),
        ('p', 'As facturas emitidas incluem um hash encadeado e uma assinatura RSA gerada pela aplicação. Estes mecanismos descrevem o funcionamento do código; não constituem, por si só, uma auditoria independente, uma garantia absoluta de segurança ou prova de conformidade fiscal.'),
        ('p', 'A aplicação usa Next.js e TypeScript, Prisma para acesso a dados e PostgreSQL como base relacional. O projecto contém configuração de Supabase, enquanto o fluxo de autenticação da aplicação é actualmente gerido por Auth.js. A disponibilidade de integrações externas depende das credenciais e da configuração de cada ambiente.'),
    ]),
]
closing = [
    ('Visão do projecto', [
        ('p', 'A visão da Machine Facturas é consolidar uma ferramenta clara para os processos comerciais e administrativos que as empresas usam no dia a dia. A evolução do produto pode continuar a partir das necessidades observadas, com melhoria gradual da experiência e dos fluxos existentes, mantendo utilidade e adequação ao contexto angolano como critérios de desenvolvimento.'),
        ('p', 'Não foi encontrada uma lista pública de funcionalidades futuras confirmadas. Por isso, esta apresentação descreve o que está implementado e assinala separadamente a área experimental, sem transformar ideias ou possibilidades em compromissos de produto.'),
    ]),
    ('Autoria e desenvolvimento', [
        ('p', 'A Machine Facturas foi concebida e desenvolvida por Joaquim Zumbuca. O trabalho traduz uma iniciativa tecnológica própria, uma visão de utilização prática do software e o compromisso com a criação de soluções digitais úteis para as empresas, especialmente no contexto angolano.'),
    ]),
    ('Conclusão', [
        ('p', 'A proposta da plataforma é reunir facturação, registos comerciais e informação financeira num fluxo organizado, com funcionalidades verificáveis no código e uma direcção de desenvolvimento centrada nas necessidades empresariais. Para conhecer o projecto e acompanhar a sua evolução, consulte o website indicado na documentação disponível.'),
    ]),
    ('Contactos e presença digital', [
        ('p', 'Website: https://facturacao-angola.vercel.app'),
        ('p', 'Telefone/WhatsApp: 937 555 618'),
        ('p', 'Email: machinefacturas@gmail.com'),
        ('p', 'Redes sociais: BlackMachine'),
    ]),
]

linkedin = '''Tenho trabalhado directamente no desenvolvimento da Machine Facturas, uma plataforma de facturação e gestão empresarial pensada para ajudar as empresas a organizar clientes, produtos, facturas, pagamentos e informação financeira num só espaço de trabalho.

O projecto reúne fluxos para preparar e emitir documentos em PDF, acompanhar pagamentos e consultar ou exportar relatórios. Está a ser desenvolvido com atenção às necessidades das empresas angolanas e à organização prática dos seus processos administrativos.

A Machine Facturas foi concebida e desenvolvida por mim, Joaquim Zumbuca, como uma iniciativa para criar soluções digitais úteis no nosso contexto empresarial. A apresentação que partilho descreve as funcionalidades confirmadas e assinala a área que ainda é experimental.

Conheça o projecto: https://facturacao-angola.vercel.app

Telefone/WhatsApp: 937 555 618 | machinefacturas@gmail.com
Redes sociais: BlackMachine

#MachineFacturas #Tecnologia #Angola'''
MD_PATH.write_text('# Publicação para LinkedIn\n\n' + linkedin + '\n', encoding='utf-8')

# ---------- DOCX ----------
def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=160, start=180, bottom=160, end=180):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tc_mar.find(qn('w:'+m))
        if node is None:
            node = OxmlElement('w:'+m)
            tc_mar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def add_page_number(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    run._r.addnext(fld)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def add_body(doc, text):
    p = doc.add_paragraph(text, style='Normal')
    p.paragraph_format.keep_together = True
    return p

def add_section(doc, title, blocks):
    h = doc.add_paragraph(title, style='Heading 1')
    h.paragraph_format.keep_with_next = True
    for kind, text in blocks:
        if kind == 'p':
            add_body(doc, text)
        elif kind == 'bullet':
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(.24)
            p.paragraph_format.first_line_indent = Inches(-.13)
            p.paragraph_format.space_after = Pt(4)
            p.add_run(text)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = Inches(.78); sec.bottom_margin = Inches(.67); sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
sec.different_first_page_header_footer = True
styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(10.3); normal.font.color.rgb = RGBColor(35, 46, 61)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.10
for name, size, before, after in [('Title', 25, 0, 6), ('Subtitle', 14, 0, 12), ('Heading 1', 18, 2, 9), ('Heading 2', 12, 7, 3)]:
    st = styles[name]; st.font.name = 'Arial'; st.font.size = Pt(size); st.font.color.rgb = RGBColor(0,0,0); st.font.bold = name != 'Subtitle'
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after); st.paragraph_format.keep_with_next = True
styles['Title'].font.bold = True
styles['Subtitle'].font.bold = False
styles['List Bullet'].font.name = 'Arial'; styles['List Bullet'].font.size = Pt(10.1)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run('MACHINE FACTURAS  |  APRESENTAÇÃO INSTITUCIONAL')
r.font.name = 'Arial'; r.font.size = Pt(8); r.font.color.rgb = RGBColor(2,31,135); r.bold = True
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('Machine Facturas  ·  Joaquim Zumbuca  ·  ')
r.font.size = Pt(8); r.font.color.rgb = RGBColor(105,117,132)
add_page_number(footer)

# Cover page: clean white editorial layout with a restrained project-blue accent.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(42); p.paragraph_format.space_after = Pt(20)
r = p.add_run('APRESENTAÇÃO INSTITUCIONAL')
r.font.name='Arial'; r.font.size=Pt(9); r.bold=True; r.font.color.rgb=RGBColor(2,31,135)
p = doc.add_paragraph(style='Title'); p.paragraph_format.space_after = Pt(5)
p.add_run('Machine Facturas')
p = doc.add_paragraph(style='Subtitle')
p.add_run('Facturação e gestão empresarial para empresas em Angola')
# Brand mark on white field, preserving original project logo proportions.
p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(20)
p.add_run().add_picture(str(LOGO), width=Inches(4.2))
p = doc.add_paragraph()
p.paragraph_format.space_after=Pt(26)
r=p.add_run('Uma plataforma digital para organizar clientes, produtos, documentos comerciais e informação financeira num só espaço de trabalho.')
r.font.name='Arial'; r.font.size=Pt(13); r.font.color.rgb=RGBColor(24,39,60)
t = doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False; t.columns[0].width=Inches(6.5)
cell=t.cell(0,0); shade_cell(cell,'EEF2FF'); set_cell_margins(cell,top=200,bottom=200,start=240,end=240)
p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2)
r=p.add_run('Desenvolvido por Joaquim Zumbuca'); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor(7,25,54)
p=cell.add_paragraph('Criador da Machine Facturas | Desenvolvedor de Software'); p.paragraph_format.space_after=Pt(0)
p.runs[0].font.size=Pt(9.5); p.runs[0].font.color.rgb=RGBColor(75,91,110)
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(28)
r=p.add_run('Website indicado na documentação do projecto: facturacao-angola.vercel.app')
r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(94,108,126)
doc.add_page_break()

add_section(doc, 'Apresentação da Machine Facturas', sections[0][1])
add_section(doc, 'O desafio', sections[1][1])
add_section(doc, 'A solução', sections[2][1])
doc.add_page_break()

h=doc.add_paragraph('Principais funcionalidades', style='Heading 1')
add_body(doc, 'As capacidades abaixo foram confirmadas no código analisado. A sua disponibilidade num ambiente publicado depende da configuração operacional do projecto.')
for title, body in features:
    h=doc.add_paragraph(title, style='Heading 2')
    add_body(doc, body)
# Experimental feature, explicitly separated from the implemented capabilities.
t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
cell=t.cell(0,0); shade_cell(cell,'F5F6F8'); set_cell_margins(cell,top=120,bottom=120,start=180,end=180)
p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(3)
r=p.add_run('Área parcial / experimental'); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(2,31,135)
p=cell.add_paragraph('Exportação XML identificada como SAFT-AO: a implementação actual é simplificada e contém valores predefinidos. Não é apresentada como exportação fiscal validada ou integração oficial com a AGT.')
p.paragraph_format.space_after=Pt(0); p.runs[0].font.size=Pt(9.2)
doc.add_page_break()

add_section(doc, 'Pensada para o contexto angolano', angola_security[0][1])
add_section(doc, 'Segurança e organização da informação', angola_security[1][1])
doc.add_page_break()

for title, blocks in closing:
    add_section(doc, title, blocks)

# Internal source note: the only confirmed web address comes from the project README.
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10)
r=p.add_run('Nota de âmbito: descrição baseada no código-fonte e na documentação do projecto disponíveis em 22 de Outubro de 2026. Contactos fornecidos pelo responsável do projecto: telefone/WhatsApp 937 555 618, email machinefacturas@gmail.com e redes sociais BlackMachine.')
r.font.size=Pt(8); r.font.italic=True; r.font.color.rgb=RGBColor(105,117,132)
doc.core_properties.title='Machine Facturas - Apresentação Institucional'
doc.core_properties.subject='Apresentação institucional e funcionalidades verificadas no código'
doc.core_properties.author='Joaquim Zumbuca'
doc.save(DOCX_PATH)

# ---------- PDF (ReportLab) ----------
FONT_REG = Path(r'C:\Windows\Fonts\arial.ttf')
FONT_BOLD = Path(r'C:\Windows\Fonts\arialbd.ttf')
pdfmetrics.registerFont(TTFont('ArialMF', str(FONT_REG)))
pdfmetrics.registerFont(TTFont('ArialMF-Bold', str(FONT_BOLD)))
NAVY=colors.HexColor('#071936'); BLUE=colors.HexColor('#021F87'); TEAL=colors.HexColor('#248468'); PALE=colors.HexColor('#F5F6F8'); INK=colors.HexColor('#18273C'); MUTED=colors.HexColor('#68798D'); LINE=colors.HexColor('#DCE2E9')
styles_pdf={
    'kicker':ParagraphStyle('kicker',fontName='ArialMF-Bold',fontSize=8.2,leading=11,textColor=BLUE,spaceAfter=4),
    'h1':ParagraphStyle('h1',fontName='ArialMF-Bold',fontSize=18,leading=22,textColor=colors.black,spaceBefore=2,spaceAfter=9,keepWithNext=True),
    'h2':ParagraphStyle('h2',fontName='ArialMF-Bold',fontSize=11.2,leading=14,textColor=colors.black,spaceBefore=5,spaceAfter=3,keepWithNext=True),
    'body':ParagraphStyle('body',fontName='ArialMF',fontSize=9.25,leading=13.3,textColor=INK,spaceAfter=6,alignment=TA_LEFT),
    'small':ParagraphStyle('small',fontName='ArialMF',fontSize=8.2,leading=11,textColor=MUTED,spaceAfter=3),
    'card':ParagraphStyle('card',fontName='ArialMF',fontSize=8.7,leading=12,textColor=INK,spaceAfter=0),
    'cardtitle':ParagraphStyle('cardtitle',fontName='ArialMF-Bold',fontSize=10.2,leading=13,textColor=NAVY,spaceAfter=3),
}
class PublicationDoc(BaseDocTemplate):
    def __init__(self, path):
        super().__init__(str(path), pagesize=A4, leftMargin=17*mm, rightMargin=17*mm, topMargin=24*mm, bottomMargin=19*mm, title='Machine Facturas - Apresentação Institucional', author='Joaquim Zumbuca', subject='Facturação e gestão empresarial para empresas em Angola')
        w,h=A4
        cover_frame=Frame(0,0,w,h,id='cover',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)
        body_frame=Frame(17*mm,19*mm,w-34*mm,h-43*mm,id='body',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)
        self.addPageTemplates([PageTemplate(id='Cover',frames=[cover_frame],onPage=self.draw_cover),PageTemplate(id='Body',frames=[body_frame],onPage=self.draw_body)])
    def draw_cover(self, canvas, doc):
        w,h=A4
        canvas.saveState(); canvas.setFillColor(NAVY); canvas.rect(0,0,w,h,fill=1,stroke=0)
        canvas.setFillColor(BLUE); canvas.rect(0,h-15*mm,w,15*mm,fill=1,stroke=0)
        canvas.setFillColor(colors.HexColor('#248468')); canvas.rect(0,h-16.5*mm,w,1.5*mm,fill=1,stroke=0)
        # low-contrast geometry echoes the application's clean interface palette.
        canvas.setStrokeColor(colors.HexColor('#243756')); canvas.setLineWidth(0.8)
        for i in range(6): canvas.circle(w+15*mm, h/2-35*mm, (35+i*11)*mm, stroke=1, fill=0)
        canvas.setFillColor(colors.HexColor('#EAF0F8')); canvas.roundRect(20*mm,h-119*mm,106*mm,38*mm,5*mm,fill=1,stroke=0)
        canvas.drawImage(str(LOGO),26*mm,h-113*mm,width=94*mm,height=25*mm,preserveAspectRatio=True,anchor='c',mask='auto')
        canvas.setFillColor(colors.HexColor('#9ED9D2')); canvas.setFont('ArialMF-Bold',8.5); canvas.drawString(20*mm,h-139*mm,'APRESENTAÇÃO INSTITUCIONAL')
        canvas.setFillColor(colors.white); canvas.setFont('ArialMF-Bold',28); canvas.drawString(20*mm,h-157*mm,'Machine Facturas')
        canvas.setFillColor(colors.HexColor('#D8E1EC')); canvas.setFont('ArialMF',15); canvas.drawString(20*mm,h-171*mm,'Facturação e gestão empresarial')
        canvas.drawString(20*mm,h-179*mm,'para empresas em Angola')
        canvas.setFillColor(colors.HexColor('#C8D3E0')); canvas.setFont('ArialMF',10)
        lines=['Uma plataforma digital para organizar clientes, produtos,', 'documentos comerciais e informação financeira num só', 'espaço de trabalho.']
        for i,line in enumerate(lines): canvas.drawString(20*mm,h-199*mm-i*6*mm,line)
        canvas.setStrokeColor(colors.HexColor('#40516A')); canvas.line(20*mm,42*mm,w-20*mm,42*mm)
        canvas.setFillColor(colors.white); canvas.setFont('ArialMF-Bold',10.5); canvas.drawString(20*mm,33*mm,'Desenvolvido por Joaquim Zumbuca')
        canvas.setFillColor(colors.HexColor('#C8D3E0')); canvas.setFont('ArialMF',9); canvas.drawString(20*mm,26*mm,'Criador da Machine Facturas | Desenvolvedor de Software')
        canvas.setFillColor(colors.HexColor('#9DAEC2')); canvas.setFont('ArialMF',8); canvas.drawString(20*mm,15*mm,'facturacao-angola.vercel.app  |  Website indicado na documentação do projecto')
        canvas.restoreState()
    def draw_body(self, canvas, doc):
        w,h=A4; canvas.saveState()
        canvas.setStrokeColor(LINE); canvas.setLineWidth(.6); canvas.line(17*mm,h-14*mm,w-17*mm,h-14*mm)
        canvas.setFont('ArialMF-Bold',7.5); canvas.setFillColor(BLUE); canvas.drawString(17*mm,h-11*mm,'MACHINE FACTURAS')
        canvas.setFont('ArialMF',7.5); canvas.setFillColor(MUTED); canvas.drawRightString(w-17*mm,h-11*mm,'APRESENTAÇÃO INSTITUCIONAL')
        canvas.line(17*mm,13*mm,w-17*mm,13*mm)
        canvas.setFont('ArialMF',7.3); canvas.setFillColor(MUTED); canvas.drawString(17*mm,8.5*mm,'Joaquim Zumbuca  |  Descrição baseada no código e documentação do projecto')
        canvas.drawRightString(w-17*mm,8.5*mm,str(doc.page))
        canvas.restoreState()

def P(text, style='body'):
    return Paragraph(text,styles_pdf[style])
def sec(title, blocks):
    out=[P(title,'h1')]
    for kind,text in blocks:
        if kind=='p': out.append(P(text))
    return out

story=[NextPageTemplate('Body'),PageBreak()]
story += [P('APRESENTAÇÃO DA PLATAFORMA','kicker'),P('Apresentação da Machine Facturas','h1')]
story += [P(t) for k,t in sections[0][1]]
story += [P('O desafio','h1')]+[P(t) for k,t in sections[1][1]]
story += [P('A solução','h1')]+[P(t) for k,t in sections[2][1]]
story += [PageBreak(),P('CAPACIDADES CONFIRMADAS NO CÓDIGO','kicker'),P('Principais funcionalidades','h1'),P('As capacidades abaixo foram confirmadas no código analisado. A sua disponibilidade num ambiente publicado depende da configuração operacional do projecto.')]
for title,body in features:
    story.append(KeepTogether([P(title,'h2'),P(body)]))
experimental=Table([[P('<b>Área parcial / experimental</b>','cardtitle')],[P('Exportação XML identificada como SAFT-AO: a implementação actual é simplificada e contém valores de exemplo predefinidos. Não é apresentada como exportação fiscal validada ou integração oficial com a AGT.','card')]],colWidths=[176*mm])
experimental.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),PALE),('BOX',(0,0),(-1,-1),.5,LINE),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
story += [Spacer(1,4),experimental]
story += [PageBreak(),P('ANGOLA E RESPONSABILIDADE','kicker')]
for title,blocks in angola_security:
    story += [P(title,'h1')]+[P(t) for k,t in blocks]
story += [PageBreak(),P('DIRECÇÃO E IDENTIFICAÇÃO','kicker')]
for title,blocks in closing:
    story += [P(title,'h1')]+[P(t) for k,t in blocks]
story += [Spacer(1,5),P('Nota de âmbito: descrição baseada no código-fonte e na documentação do projecto disponíveis em 22 de Outubro de 2026. Contactos fornecidos pelo responsável do projecto: telefone/WhatsApp 937 555 618, email machinefacturas@gmail.com e redes sociais BlackMachine.','small')]

pdf=PublicationDoc(PDF_PATH)
pdf.build(story)
print(f'Created: {DOCX_PATH.name}, {PDF_PATH.name}, {MD_PATH.name}')